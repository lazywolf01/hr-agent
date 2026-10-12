"""
HR Candidate Selection Agent
----------------------------
Pipeline:
  1. Job description (plain text)  -> LLM -> JobDescription JSON
  2. Every resume (.pdf / .docx)   -> text -> LLM -> Resume JSON
  3. JD JSON + each Resume JSON    -> LLM -> CandidateEvaluation JSON (+ suitability %)
  4. Rank all candidates, send top 2 to LLM -> why they are top + their top skills
  5. Everything is saved as JSON in the output folder

Usage:
  python hr_agent.py --jd job_description.txt --resumes resume --out output
"""
import argparse
import json
import os
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field, ValidationError
from pypdf import PdfReader

load_dotenv()
my_key = os.getenv("GROQ_API_KEY")
if not my_key:
    raise ValueError("API ERROR: GROQ_API_KEY not found in .env")

client = Groq(api_key=my_key)
model = "openai/gpt-oss-120b"

# Weights used to compute the final suitability % (must add up to 1.0)
WEIGHTS = {
    "skills_match": 0.40,
    "experience_match": 0.30,
    "education_match": 0.15,
    "projects_and_certifications_match": 0.15,
}


# ======================= SCHEMAS =======================

# ---------- Job description ----------
class ExperienceRequirement(BaseModel):
    min_years: float | None = Field(None, description="Minimum years of experience required")
    max_years: float | None = Field(None, description="Maximum years of experience, if mentioned")
    domain: str | None = Field(None, description="Domain/industry the experience should be in")


class JobDescription(BaseModel):
    job_title: str
    department: str | None = None
    location: str | None = None
    employment_type: str | None = Field(None, description="Full-time, Part-time, Contract, Internship ...")
    work_mode: str | None = Field(None, description="Onsite, Remote or Hybrid")
    experience: ExperienceRequirement
    education_requirements: list[str] = Field(default_factory=list)
    required_skills: list[str] = Field(default_factory=list, description="Must-have skills")
    preferred_skills: list[str] = Field(default_factory=list, description="Nice-to-have skills")
    responsibilities: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    salary_range: str | None = None


# ---------- Resume ----------
class WorkExperience(BaseModel):
    company: str
    job_title: str
    start_date: str | None = None
    end_date: str | None = Field(None, description="'Present' if currently working")
    duration_years: float | None = None
    highlights: list[str] = Field(default_factory=list)
    technologies: list[str] = Field(default_factory=list)


class Education(BaseModel):
    degree: str
    field_of_study: str | None = None
    institution: str | None = None
    graduation_year: str | None = None
    grade: str | None = Field(None, description="CGPA / percentage, if given")


class Project(BaseModel):
    name: str
    description: str | None = None
    technologies: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    full_name: str
    email: str | None = None
    phone: str | None = None
    location: str | None = None
    linkedin: str | None = None
    github: str | None = None
    summary: str | None = None
    total_experience_years: float = Field(0, description="Total professional experience in years")
    skills: list[str] = Field(default_factory=list)
    work_experience: list[WorkExperience] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    languages: list[str] = Field(default_factory=list)


# ---------- Evaluation ----------
class ScoreBreakdown(BaseModel):
    skills_match: int = Field(..., ge=0, le=100)
    experience_match: int = Field(..., ge=0, le=100)
    education_match: int = Field(..., ge=0, le=100)
    projects_and_certifications_match: int = Field(..., ge=0, le=100)


class CandidateEvaluation(BaseModel):
    candidate_name: str
    score_breakdown: ScoreBreakdown
    matched_required_skills: list[str]
    missing_required_skills: list[str]
    matched_preferred_skills: list[str]
    strengths: list[str]
    gaps: list[str]
    recommendation: str = Field(..., description="One of: Strong Fit, Good Fit, Partial Fit, Not a Fit")
    reasoning: str = Field(..., description="2-4 sentence justification")


# ---------- Top candidates ----------
class TopCandidate(BaseModel):
    rank: int
    candidate_name: str
    suitability_percentage: float
    top_skills: list[str] = Field(..., description="3-5 skills that most made this candidate stand out for this job")
    why_selected: str


class TopCandidatesResult(BaseModel):
    top_candidates: list[TopCandidate]
    comparison_summary: str = Field(..., description="How the top candidates compare to each other")


# ======================= LLM HELPER =======================

def ask_llm(system_prompt: str, user_prompt: str, schema: type[BaseModel], retries: int = 2) -> BaseModel:
    """Call the LLM, force JSON output matching `schema`, validate with pydantic (retry on failure)."""
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]
    response_format = {
        "type": "json_schema",
        "json_schema": {"name": schema.__name__, "schema": schema.model_json_schema()},
    }
    last_error = None
    for _ in range(retries + 1):
        response = client.chat.completions.create(
            model=model,
            messages=messages,
            temperature=0,
            response_format=response_format,
        )
        content = response.choices[0].message.content
        try:
            return schema.model_validate_json(content)
        except ValidationError as e:
            last_error = e
            # Tell the model what went wrong and let it fix it
            messages += [
                {"role": "assistant", "content": content},
                {"role": "user", "content": f"Your JSON did not match the schema:\n{e}\nReturn corrected JSON only."},
            ]
    raise RuntimeError(f"LLM output never matched {schema.__name__}: {last_error}")


# ======================= STEP 1: JOB DESCRIPTION =======================

def parse_job_description(jd_text: str) -> JobDescription:
    system_prompt = (
        "You are an HR assistant. Convert the job description into structured JSON. "
        "Only use information present in the text; use null / empty lists for anything not mentioned. "
        "Split skills into required (must-have) and preferred (nice-to-have). Keep skill names short, e.g. 'Python', 'Docker'."
    )
    return ask_llm(system_prompt, f"Job description:\n{jd_text}", JobDescription)


# ======================= STEP 2: RESUMES =======================

def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if path.suffix.lower() == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        # Many resumes put content inside tables
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def parse_resume(resume_text: str) -> Resume:
    system_prompt = (
        "You are a resume parser. Convert the resume text into structured JSON. "
        "Only use information present in the resume; never invent details. "
        "Calculate total_experience_years from work experience dates (internships count as their actual duration). "
        "Keep skill names short and normalized, e.g. 'Python', 'Machine Learning', 'SQL'."
    )
    return ask_llm(system_prompt, f"Resume text:\n{resume_text}", Resume)


# ======================= STEP 3: EVALUATION =======================

def evaluate_candidate(jd: JobDescription, resume: Resume) -> CandidateEvaluation:
    system_prompt = (
        "You are an expert technical recruiter. Compare the candidate resume against the job description, both given as JSON. "
        "Score each category from 0 to 100:\n"
        "- skills_match: coverage of required skills (most important) and preferred skills, counting equivalent technologies\n"
        "- experience_match: years and relevance of experience versus the requirement\n"
        "- education_match: degree / field versus the requirement (100 if no requirement is stated and candidate has relevant education)\n"
        "- projects_and_certifications_match: how relevant their projects and certifications are to the responsibilities\n"
        "Be strict and objective. Base everything only on the given data. "
        "Refer to the candidate by name or as 'they'; never assume gender."
    )
    user_prompt = (
        f"JOB DESCRIPTION:\n{jd.model_dump_json(indent=2)}\n\n"
        f"CANDIDATE RESUME:\n{resume.model_dump_json(indent=2)}"
    )
    return ask_llm(system_prompt, user_prompt, CandidateEvaluation)


def suitability_percentage(scores: ScoreBreakdown) -> float:
    """Weighted average of the category scores -> same formula for every candidate."""
    return round(sum(getattr(scores, k) * w for k, w in WEIGHTS.items()), 1)


# ======================= STEP 4: TOP CANDIDATES =======================

def pick_top_candidates(jd: JobDescription, ranked: list[dict], top_n: int = 2) -> TopCandidatesResult:
    top = ranked[:top_n]
    system_prompt = (
        f"You are a hiring manager. These are the top {len(top)} candidates for the job, already ranked by suitability. "
        "Keep the given rank and suitability_percentage. For each one, list the 3-5 skills that most made them stand out "
        "for THIS job and explain in 2-3 sentences why they were selected. Then compare them briefly. "
        "Refer to candidates by name or as 'they'; never assume gender."
    )
    user_prompt = (
        f"JOB DESCRIPTION:\n{jd.model_dump_json(indent=2)}\n\n"
        f"TOP CANDIDATES:\n{json.dumps(top, indent=2)}"
    )
    return ask_llm(system_prompt, user_prompt, TopCandidatesResult)


# ======================= MAIN PIPELINE =======================

def save_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def run(jd_path: Path, resume_dir: Path, out_dir: Path, top_n: int = 2) -> dict:
    # Step 1
    print("Step 1: Parsing job description...")
    jd = parse_job_description(jd_path.read_text(encoding="utf-8"))
    save_json(out_dir / "job_description.json", jd.model_dump())

    # Step 2
    resume_files = sorted(p for p in resume_dir.iterdir() if p.suffix.lower() in (".pdf", ".docx"))
    if not resume_files:
        raise FileNotFoundError(f"No .pdf or .docx resumes found in {resume_dir}")
    print(f"Step 2: Parsing {len(resume_files)} resume(s)...")

    resumes: list[tuple[str, Resume]] = []
    for file in resume_files:
        try:
            text = extract_text(file)
            if not text.strip():
                print(f"  ! {file.name}: no text found (scanned image?) - skipped")
                continue
            resume = parse_resume(text)
            save_json(out_dir / "parsed_resumes" / f"{file.stem}.json", resume.model_dump())
            resumes.append((file.name, resume))
            print(f"  OK {file.name} -> {resume.full_name}")
        except Exception as e:
            print(f"  ! {file.name}: failed ({e}) - skipped")

    # Step 3
    print("Step 3: Evaluating candidates against the job...")
    results = []
    for file_name, resume in resumes:
        try:
            ev = evaluate_candidate(jd, resume)
        except Exception as e:
            print(f"  ! {file_name}: evaluation failed ({e}) - skipped")
            continue
        score = suitability_percentage(ev.score_breakdown)
        results.append({
            "file_name": file_name,
            "suitability_percentage": score,
            **ev.model_dump(),
        })
        print(f"  {resume.full_name}: {score}%")

    ranked = sorted(results, key=lambda r: r["suitability_percentage"], reverse=True)
    for i, r in enumerate(ranked, start=1):
        r["rank"] = i

    # Step 4
    print(f"Step 4: Selecting top {top_n} candidates...")
    top = pick_top_candidates(jd, ranked, top_n) if ranked else TopCandidatesResult(
        top_candidates=[], comparison_summary="No candidates could be evaluated."
    )

    report = {
        "job_title": jd.job_title,
        "total_resumes": len(resume_files),
        "candidates_evaluated": len(ranked),
        "scoring_weights": WEIGHTS,
        "top_candidates": top.model_dump()["top_candidates"],
        "comparison_summary": top.comparison_summary,
        "all_candidates_ranked": ranked,
    }
    save_json(out_dir / "final_report.json", report)
    return report


def print_summary(report: dict) -> None:
    print("\n" + "=" * 60)
    print(f"RESULTS FOR: {report['job_title']}")
    print("=" * 60)
    for r in report["all_candidates_ranked"]:
        print(f"#{r['rank']:<3} {r['candidate_name']:<30} {r['suitability_percentage']:>5}%  {r['recommendation']}")
    print("\nTOP CANDIDATES")
    for t in report["top_candidates"]:
        print(f"\n#{t['rank']} {t['candidate_name']} ({t['suitability_percentage']}%)")
        print(f"   Top skills : {', '.join(t['top_skills'])}")
        print(f"   Why        : {t['why_selected']}")
    print(f"\n{report['comparison_summary']}")


if __name__ == "__main__":
    here = Path(__file__).parent
    parser = argparse.ArgumentParser(description="HR agent: rank resumes against a job description")
    parser.add_argument("--jd", type=Path, default=here / "job_description.txt", help="Plain-text job description file")
    parser.add_argument("--resumes", type=Path, default=here / "resume", help="Folder with .pdf / .docx resumes")
    parser.add_argument("--out", type=Path, default=here / "output", help="Folder to write JSON results")
    parser.add_argument("--top", type=int, default=2, help="How many top candidates to pick")
    args = parser.parse_args()

    report = run(args.jd, args.resumes, args.out, args.top)
    print_summary(report)
    print(f"\nFull JSON saved in: {args.out}")
